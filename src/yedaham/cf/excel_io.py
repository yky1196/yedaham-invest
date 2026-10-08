"""엑셀 모델(Ver2_CF_모델.xlsx) ↔ 엔진 입력/정답값 변환.

- `load_inputs(path)` : 수식 원본 파일에서 Back Data를 읽어 ModelInputs 생성
- `load_golden(path)` : LibreOffice로 재계산한 파일에서 엑셀 결과값(정답) 추출

엑셀의 행 순서 규칙(04·05·07·10·11·12가 같은 행)을 여기서 한 번만 해석하고,
엔진 안에서는 자산 이름으로만 연결한다.
"""
from __future__ import annotations

import re
from typing import Any, Optional

import openpyxl

from .model import (Asset, Assumptions, BondLine, Curves, DivOverride, FinParams,
                    FinProduct, FXPath, ModelInputs, T1Event)

S_ASSUM = "01_가정"
S_FX = "02_BD_환율"
S_CURVE = "03_BD_곡선"
S_MASTER = "04_BD_자산마스터"
S_CAL = "05_BD_캘린더"
S_T1 = "06_BD_T1확정"
S_PIPE = "07_BD_파이프라인"
S_FIN = "08_BD_금융상품"
S_BOND = "13_채권별"


def _num(v: Any, default: float = 0.0) -> float:
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.startswith("="):
        raise ValueError(f"입력 셀에 수식이 있습니다: {v[:60]}")
    return float(v)


def _opt_num(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    return _num(v)


def _opt_str(v: Any) -> Optional[str]:
    if v is None or v == "":
        return None
    return str(v)


def _col(ws, letter: str, r: int):
    return ws[f"{letter}{r}"].value


def _curve_block(ws, first_row: int, last_row: int) -> dict[str, list[Optional[float]]]:
    out = {}
    for r in range(first_row, last_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        vals = []
        for c in range(2, 15):  # B..N = Y0..Y12
            v = ws.cell(r, c).value
            vals.append(float(v) if isinstance(v, (int, float)) else None)
        out[str(name)] = vals
    return out


RE_DIV_YIELD = re.compile(r"^=IF\(\$AX\d+=\"Matured\",0,\$L\d+\*([0-9.]+)\)\*'01_가정'!\$B\$10$")
RE_DIV_STD = re.compile(r"^=IF\(\$AX(\d+)=\"Matured\",0,IF\(\$V\1=\"Y\"")
RE_ROUND = re.compile(r"^=ROUND\(([-0-9.]+),4\)\*\$C\d+$")
RE_PACE_RESID = re.compile(r"^=MAX\(0,\(1-SUM\(\$B\d+:[A-Z]+\d+\)\)\*IFERROR\(INDEX\('03_BD_곡선'!\$B\$3:\$N\$15,MATCH\(\"([^\"]+)\",'03_BD_곡선'!\$A\$3:\$A\$15,0\),(\d+)\),0\)/12\)$")


def load_inputs(path: str) -> ModelInputs:
    wb = openpyxl.load_workbook(path)  # 수식 그대로
    warnings: list[str] = []

    # 01 가정
    ws = wb[S_ASSUM]
    b = lambda r: _num(ws[f"B{r}"].value)  # noqa: E731
    assum = Assumptions(
        opening_liquidity=b(4), fin_yield=b(5), blind_draw_cap=b(6), project_draw_cap=b(7),
        fx_scen_usd=b(8), fx_scen_eur=b(9), div_mult=b(10), repay_mult=b(11), call_mult=b(12),
        plan_switch=b(13), bond_wht=b(14), liquidation_months=b(15), reinvest_decay=b(16),
    )

    # 02 환율
    ws = wb[S_FX]
    months, usd, eur = [], [], []
    for r in range(5, 22):
        months.append(str(ws.cell(r, 1).value))
        usd.append(_num(ws.cell(r, 2).value))
        eur.append(_num(ws.cell(r, 3).value))
    fx = FXPath(months, usd, eur, spot_usd=_num(ws["D2"].value), spot_eur=_num(ws["E2"].value))
    n = len(months)

    # 03 곡선
    ws = wb[S_CURVE]
    curves = Curves(draw=_curve_block(ws, 3, 15), runoff=_curve_block(ws, 19, 31),
                    dividend=_curve_block(ws, 35, 47))

    # 04 자산마스터 + 05 캘린더 + 07 파이프라인 (같은 행)
    wm, wc, wp = wb[S_MASTER], wb[S_CAL], wb[S_PIPE]
    assets: list[Asset] = []
    for r in range(2, wm.max_row + 1):
        name = _col(wm, "A", r)
        if not name:
            continue
        if wc.cell(r, 1).value != name or wp.cell(r, 1).value != name:
            raise ValueError(f"행 순서 불일치 r={r}: {name}")
        g = lambda L: _col(wm, L, r)  # noqa: E731
        a = Asset(
            name=str(name), kind=str(g("B")), structure=str(g("C")), midclass=str(g("D") or ""),
            strategy=str(g("E") or ""), sector=str(g("F") or ""), region=str(g("G") or ""),
            ccy=str(g("H")), vintage=int(g("I")) if g("I") not in (None, "") else None,
            age=_opt_num(g("J")), commitment=_num(g("K")), invested_bal=_num(g("L")),
            unfunded=_num(g("M")), cum_called=_num(g("N")), call_12m=_num(g("P")),
            div_12m=_num(g("Q")), div_24m_ann=_num(g("R")), repay_36m=_num(g("S")),
            repay_rate_3y=_num(g("T")), recurrence=_opt_str(g("U")), sale_recovery=_opt_num(g("W")),
            call_basis=_opt_str(g("AP")), div_basis=_opt_str(g("AQ")),
            trust_end_months=_opt_num(g("AU")), rec_div_12m=_num(g("AV")), rec_div_24m=_num(g("AW")),
            status=str(g("AX")), k3_invalid=_opt_num(g("BN")), k3_invalid_note=_opt_str(g("BO")),
            invest_end=_opt_str(g("BP")), confirmed_unfunded=_opt_num(g("BS")),
            unfunded_note=_opt_str(g("BT")), win_invest_months=_num(g("BV")),
            win_post_months=_num(g("BW")),
        )
        # 연간 배당 수식이 표준과 다른 행 → 자산별 지정으로 변환
        z = g("Z")
        if isinstance(z, str) and z.startswith("="):
            m = RE_DIV_YIELD.match(z)
            if m:
                a.div_override = DivOverride("yield", float(m.group(1)), "엑셀 Z열 수동 수익률")
            elif not RE_DIV_STD.match(z):
                raise ValueError(f"알 수 없는 연간 배당 수식 r={r}: {z}")
        else:
            a.div_override = DivOverride("fixed", _num(z), "엑셀 Z열 고정값")
        ag = g("AG")
        if not (isinstance(ag, str) and ag.startswith("=")):
            a.div_months_override = _num(ag)
        # 05 캘린더
        a.cal_call = [int(_num(wc.cell(r, c).value)) for c in range(2, 14)]
        a.cal_div = [int(_num(wc.cell(r, c).value)) for c in range(14, 26)]
        a.cal_repay = [int(_num(wc.cell(r, c).value)) for c in range(26, 38)]
        a.cal_sale = [int(_num(wc.cell(r, c).value)) for c in range(38, 38 + n)]
        # 07 파이프라인
        pace, resid_idx, resid_mid, resid_year = [], [], None, 0
        for k, c in enumerate(range(2, 2 + n)):
            v = wp.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                m = RE_PACE_RESID.match(v)
                if not m:
                    raise ValueError(f"알 수 없는 페이싱 수식 r={r}: {v}")
                resid_idx.append(k)
                resid_mid, resid_year = m.group(1), int(m.group(2)) - 1
                pace.append(0.0)
            else:
                pace.append(_num(v))
        a.pipe_pacing = pace
        if resid_idx:
            a.pipe_pacing_rule = {"type": "residual_curve", "midclass": resid_mid,
                                  "year": resid_year, "months_idx": resid_idx}
        a.pipe_div_start = _opt_str(wp.cell(r, 19).value)
        a.pipe_div_rate = _num(wp.cell(r, 20).value)
        a.pipe_div_cal = [int(_num(wp.cell(r, c).value)) for c in range(21, 33)]
        a.pipe_catchup = _num(wp.cell(r, 34).value)
        assets.append(a)

    # 06 T1
    ws = wb[S_T1]
    t1 = []
    for r in range(2, 301):
        m, nm, kd, amt = (ws.cell(r, c).value for c in range(1, 5))
        if not m or not nm:
            continue
        t1.append(T1Event(str(m), str(nm), str(kd), _num(amt), str(ws.cell(r, 5).value or ""),
                          str(ws.cell(r, 6).value or "")))

    # 08 금융상품
    ws = wb[S_FIN]
    fins = []
    for r in range(2, 11):
        nm = ws.cell(r, 1).value
        if not nm:
            continue
        fins.append(FinProduct(str(nm), str(ws.cell(r, 2).value), _num(ws.cell(r, 3).value),
                               _num(ws.cell(r, 4).value), _num(ws.cell(r, 5).value),
                               str(ws.cell(r, 6).value or "")))
    fin_params = FinParams(rate=_num(ws["B12"].value), scheduled_principal=_num(ws["B13"].value))

    # 13 채권
    ws = wb[S_BOND]
    bonds = []
    for r in range(2, ws.max_row + 1):
        nm, kd = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not nm or kd not in ("이자", "원금"):
            continue
        amts = []
        for c in range(5, 5 + n):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                m = RE_ROUND.match(v)
                if not m:
                    raise ValueError(f"알 수 없는 채권 수식 r={r}: {v}")
                amts.append(float(m.group(1)))
            else:
                amts.append(_num(v))
        bonds.append(BondLine(str(nm), str(kd), _num(ws.cell(r, 3).value, 1.0),
                              str(ws.cell(r, 4).value or ""), amts))

    return ModelInputs(assum, fx, curves, assets, t1, fins, fin_params, bonds)


def load_golden(recalc_path: str) -> dict:
    """LibreOffice로 재계산한 엑셀에서 정답값을 뽑는다 (자산별 월별 + 월별 CF)."""
    wb = openpyxl.load_workbook(recalc_path, data_only=True)
    out: dict = {"assets": {}, "monthly": []}
    for key, sheet in (("call", "10_자산별_콜"), ("div", "11_자산별_배당"), ("repay", "12_자산별_상환")):
        ws = wb[sheet]
        for r in range(2, 51):
            nm = ws.cell(r, 1).value
            if not nm:
                continue
            vals = [float(ws.cell(r, c).value or 0) for c in range(2, 19)]
            out["assets"].setdefault(str(nm), {})[key] = vals
    wm = wb[S_MASTER]
    cols = {"AM": "win_budget_call", "AN": "win_budget_div", "AO": "win_budget_repay",
            "Y": "annual_call", "Z": "annual_div", "AE": "annual_repay"}
    for r in range(2, 51):
        nm = wm.cell(r, 1).value
        if nm:
            out["assets"].setdefault(str(nm), {})["derived"] = {
                v: float(wm[f"{k}{r}"].value or 0) for k, v in cols.items()}
    ws = wb["20_월별CF"]
    for r in range(2, 19):
        row = [ws.cell(r, c).value for c in range(1, 11)]
        out["monthly"].append({
            "month": row[0], "opening": row[1], "call": row[2], "div": row[3], "repay": row[4],
            "bond_int": row[5], "bond_prin": row[6], "fin_int": row[7], "net": row[8], "closing": row[9]})
    return out
